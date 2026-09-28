"""Build the mini-report DOCX from verified repository evidence."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "mini-report-backend.docx"
FIGURES = ROOT / "docs" / "figures"
BRANCH = "https://github.com/Andrey787878/hackhaton/tree/codex/backend-mini-report"


def font(run, *, size=14, bold=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is not None:
        rfonts.set(qn("w:eastAsia"), name)
    return run


def text(doc, value, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, after=0):
    p = doc.add_paragraph(style="Normal")
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    p.paragraph_format.space_after = Pt(after)
    font(p.add_run(value))
    return p


def heading(doc, value):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(value), bold=True)
    return p


def hyperlink(paragraph, label, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    properties.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "28")
    properties.append(sz)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    element.append(run)
    paragraph._p.append(element)


def picture(doc, filename, number, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Cm(15.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    font(cap.add_run(f"Рисунок {number} – {caption}"), size=14)


def page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instruction, separate, value, end):
        run._r.append(item)
    font(run, size=12)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_number(footer)

    for label in ("ФГБОУ ВО «ЮГОРСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ»",
                  "ИНЖЕНЕРНАЯ ШКОЛА ЦИФРОВЫХ ТЕХНОЛОГИЙ"):
        p = text(doc, label, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        p.paragraph_format.space_after = Pt(7)
    for _ in range(6):
        doc.add_paragraph()
    p = text(doc, "МИНИ-ОТЧЁТ", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p.paragraph_format.space_after = Pt(6)
    font(p.runs[0], bold=True)
    text(doc, "по разработке серверной части веб-приложения",
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p = text(doc, "«Платформа онлайн-хакатонов и соревнований»",
             align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p.paragraph_format.space_after = Pt(18)
    for _ in range(4):
        doc.add_paragraph()
    text(doc, "Выполнили: студенты группы ПИ41Б, команда № 3",
         align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    text(doc, "Полуянов, Бойдаченко, Абдурахман, Борисов",
         align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    text(doc, "Проверил: Исмагулов Милан Ерикович",
         align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    for _ in range(5):
        doc.add_paragraph()
    text(doc, "г. Ханты-Мансийск, 2026 г.", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    doc.add_page_break()

    heading(doc, "1 ЦЕЛЬ И РЕЗУЛЬТАТ РАБОТЫ")
    text(doc, "Цель работы — реализовать работоспособный бэкенд платформы онлайн-хакатонов "
         "с постоянным хранением, контролем доступа, серверным дедлайном сдачи и API "
         "для команд, организаторов и экспертов.")
    text(doc, "Разработаны авторизация и роли, управление хакатонами и кейсами, команды "
         "и приглашения, приём ссылок и архивов, экспертная оценка по критериям, "
         "публичный рейтинг и журнал уведомлений. API доступен через Swagger UI.")

    heading(doc, "2 ССЫЛКА НА АКТУАЛЬНЫЙ КОД")
    text(doc, "Код серверной части опубликован в открытом репозитории GitHub "
         "в ветке codex/backend-mini-report. Адрес ветки:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    hyperlink(p, BRANCH, BRANCH)
    text(doc, "В ветке размещены исходный код, миграции БД, Docker Compose, тесты "
         "и инструкция запуска. Демо-пароли и локальная конфигурация исключены из Git.")

    heading(doc, "3 ВЫБРАННЫЙ СТЕК ТЕХНОЛОГИЙ")
    text(doc, "Python 3.11 и FastAPI выбраны для типизированного HTTP API, автоматической "
         "валидации входных данных и генерации OpenAPI/Swagger. Pydantic описывает "
         "схемы запросов и ответов.")
    text(doc, "PostgreSQL 17 хранит пользователей, права, хакатоны, команды, метаданные "
         "сдач и оценки. SQLAlchemy 2 отделяет модели данных от маршрутов, Alembic "
         "обеспечивает версионируемые миграции.")
    text(doc, "MinIO предоставляет S3-совместимое приватное хранение архивов; в SQL "
         "находятся только метаданные и ключ объекта. Docker Compose запускает API, "
         "БД, MinIO и служебные контейнеры миграции и инициализации bucket. "
         "PyJWT и bcrypt применяются для сессий и хэширования паролей; Pytest "
         "проверяет бизнес-сценарии с реальными PostgreSQL и MinIO.")

    heading(doc, "4 АРХИТЕКТУРА И СТРУКТУРА")
    text(doc, "HTTP-запрос проходит через маршрут-контроллер, проверку пользователя "
         "и входных данных, сервис бизнес-логики, ORM и PostgreSQL. Архив передаётся "
         "через S3-адаптер в MinIO. Основные каталоги представлены ниже.")
    tree = (
        "hackhaton/\n"
        "├── app/\n"
        "│   ├── main.py, config.py, db.py\n"
        "│   ├── models.py, dependencies.py\n"
        "│   ├── auth_routes.py, auth_service.py\n"
        "│   ├── rbac_routes.py, rbac_service.py\n"
        "│   ├── hackathon_routes.py, hackathon_service.py\n"
        "│   ├── team_routes.py, team_service.py\n"
        "│   ├── submission_routes.py, submission_service.py\n"
        "│   ├── scoring_routes.py, scoring_service.py\n"
        "│   ├── notification_routes.py, notification_service.py\n"
        "│   ├── *_schemas.py, storage.py, storage_bootstrap.py\n"
        "├── alembic/versions/      миграции PostgreSQL\n"
        "├── tests/                 интеграционные тесты\n"
        "├── scripts/               настройка и демосценарий\n"
        "├── Dockerfile, Dockerfile.minio\n"
        "├── docker-compose.yml, compose.test.yml\n"
        "└── requirements.txt       зависимости Python"
    )
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    font(p.add_run(tree), size=10.5, name="Courier New")
    text(doc, "Файлы *_routes.py выполняют роль контроллеров: принимают HTTP-запросы "
         "и возвращают ответы. Файлы *_service.py содержат правила доступа, "
         "дедлайны и транзакции. В models.py описаны ORM-модели, в *_schemas.py — "
         "контракты API. Каталог alembic/versions хранит историю схемы БД, "
         "tests/ — автоматические проверки.")

    doc.add_page_break()
    heading(doc, "5 ПОДТВЕРЖДЕНИЕ РАБОТОСПОСОБНОСТИ")
    text(doc, "На рисунке 1 показан Swagger UI запущенного сервера с маршрутами "
         "авторизации и управления доступом. OpenAPI доступен по адресу /openapi.json.")
    picture(doc, "swagger.png", 1, "Интерфейс Swagger UI серверного API")
    text(doc, "На рисунке 2 приведён успешный ответ GET /api/v1/health: "
         "сервер и подключение к PostgreSQL исправны.")
    picture(doc, "health.png", 2, "Ответ проверки API и базы данных")
    p = text(doc, "На рисунке 3 представлен ответ публичного GET-запроса рейтинга "
             "для кейса № 3: команда получила итог 85,00 балла по весовым критериям.")
    p.paragraph_format.keep_with_next = True
    picture(doc, "leaderboard.png", 3, "Публичный рейтинг по результатам экспертной оценки")
    text(doc, "Дополнительно выполнены 87 интеграционных тестов. Сквозной сценарий "
         "проверил публикацию хакатона, создание команды, сдачу архива, назначение "
         "эксперта, оценку и рейтинг. Сдача после дедлайна отклонена с HTTP 409; "
         "анонимное скачивание архива — с HTTP 401. Проверка alembic check не "
         "обнаружила расхождений между моделями и миграциями.")

    heading(doc, "6 ЗАКЛЮЧЕНИЕ")
    text(doc, "Основная серверная бизнес-логика платформы реализована и доступна "
         "через версионированный API /api/v1. Запуск воспроизводится по README с "
         "помощью Docker Compose. Подтверждены работа API, сохранение данных и "
         "архивов, разграничение прав, соблюдение дедлайна и расчёт рейтинга.")
    heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for item in (
        "1. Задание 3. Платформа для онлайн-хакатонов и соревнований. URL: https://yusu-source.github.io/GH_pages/03_hackathon.html (дата обращения: 28.09.2026).",
        "2. ГОСТ 7.32—2017. Отчёт о научно-исследовательской работе. Структура и правила оформления.",
        "3. Репозиторий проекта: https://github.com/Andrey787878/hackhaton.",
    ):
        text(doc, item, indent=False)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
